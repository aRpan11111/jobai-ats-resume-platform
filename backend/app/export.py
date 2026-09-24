import io
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_docx(parsed_resume) -> io.BytesIO:
    """Generate a clean, single-column ATS-friendly DOCX from structured JSON."""
    if isinstance(parsed_resume, str):
        import json
        try:
            parsed_resume = json.loads(parsed_resume)
        except:
            parsed_resume = {}
            
    document = docx.Document()
    
    # 1. Contact / Header
    contact_info = parsed_resume.get("contact_info", {})
    name = contact_info.get("name") or "Candidate Name"
    email = contact_info.get("email") or "email@example.com"
    phone = contact_info.get("phone") or "555-0100"
    linkedin = contact_info.get("linkedin") or ""
    github = contact_info.get("github") or ""
    
    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(name)
    run.bold = True
    run.font.size = Pt(14)
    
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_links = " | ".join(filter(None, [email, phone, linkedin, github]))
    contact.add_run(contact_links)

    # 2. Summary
    summary_text = parsed_resume.get("summary")
    if summary_text:
        document.add_heading("Summary", level=1)
        document.add_paragraph(summary_text)

    # 3. Skills
    skills = parsed_resume.get("skills", [])
    if skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(skills))

    # 4. Experience
    experience = parsed_resume.get("experience", [])
    if experience:
        document.add_heading("Experience", level=1)
        for exp in experience:
            p = document.add_paragraph()
            p.add_run(exp.get("title", "Title")).bold = True
            p.add_run(f" | {exp.get('company', 'Company')} | {exp.get('dates', 'Dates')}")
            for bullet in exp.get("bullets", []):
                document.add_paragraph(bullet, style='List Bullet')

    # 5. Education
    education = parsed_resume.get("education", [])
    if education:
        document.add_heading("Education", level=1)
        for edu in education:
            p = document.add_paragraph()
            p.add_run(edu.get("degree", "Degree")).bold = True
            p.add_run(f" | {edu.get('institution', 'Institution')} | {edu.get('dates', 'Dates')}")

    # 6. Projects
    projects = parsed_resume.get("projects", [])
    if projects:
        document.add_heading("Projects", level=1)
        for proj in projects:
            p = document.add_paragraph()
            p.add_run(proj.get("name", "Project Name")).bold = True
            for bullet in proj.get("bullets", []):
                document.add_paragraph(bullet, style='List Bullet')
                
    # 7. Certifications
    certifications = parsed_resume.get("certifications", [])
    if certifications:
        document.add_heading("Certifications", level=1)
        for cert in certifications:
            document.add_paragraph(cert, style='List Bullet')
            
    # 8. Achievements
    achievements = parsed_resume.get("achievements", [])
    if achievements:
        document.add_heading("Achievements", level=1)
        for ach in achievements:
            document.add_paragraph(ach, style='List Bullet')
                
    doc_io = io.BytesIO()
    document.save(doc_io)
    doc_io.seek(0)
    return doc_io
